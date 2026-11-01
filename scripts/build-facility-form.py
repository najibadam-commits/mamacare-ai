#!/usr/bin/env python3
"""Builds the one-page facility details and consent form (Word). Needs: pip install python-docx

    python3 scripts/build-facility-form.py MamaCare_AI_facility_details_form.docx
"""
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

out = sys.argv[1] if len(sys.argv) > 1 else 'MamaCare_AI_facility_details_form.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8); sec.top_margin = Cm(1.3); sec.bottom_margin = Cm(1.3)
st = doc.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
NAVY = RGBColor(0x0B, 0x3C, 0x5D)

def shade(cell, color):
    pr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color); pr.append(sh)

def para(text, bold=False, size=10, color=None, italic=False, after=4):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.space_before = Pt(0); return p

def cell_text(cell, text, bold=False, size=9.5):
    cell.text = ''; r = cell.paragraphs[0].add_run(text); r.bold = bold; r.font.size = Pt(size)
    cell.paragraphs[0].paragraph_format.space_after = Pt(1)

def table(rows, w=(6.2, 11.2), h=0.85):
    t = doc.add_table(rows=len(rows), cols=2); t.style = 'Table Grid'; t.autofit = False
    for i, (a, b) in enumerate(rows):
        cell_text(t.rows[i].cells[0], a, bold=True, size=9); shade(t.rows[i].cells[0], 'E8EEF3')
        cell_text(t.rows[i].cells[1], b); t.rows[i].height = Cm(h)
        for c, x in zip(t.rows[i].cells, w): c.width = Cm(x)
    return t

p = sec.footer.paragraphs[0]; p.text = 'MamaCare AI facility details form  |  Only provide information about the facility. Never write anything about a patient.'
p.runs[0].font.size = Pt(8); p.alignment = WD_ALIGN_PARAGRAPH.CENTER

para('MamaCare AI', bold=True, size=18, color=NAVY, after=0)
para('Health facility details and permission form', bold=True, size=13, after=4)
para('MamaCare AI is a free, open-source phone app for mothers. Its facility directory helps a woman find a clinic and call it. We only list a '
     'facility that has agreed, and we only publish details that someone at the facility has confirmed. Being listed is not an endorsement or a '
     'partnership. One form per facility, please print clearly.', size=9.5, after=6)

para('1. Facility', bold=True, size=10.5, color=NAVY, after=2)
table([('Official facility name', ''),
       ('Type (tick one)', '☐ Primary Health Centre   ☐ Maternity / Comprehensive Health Centre   ☐ General Hospital   ☐ Tertiary / Teaching   ☐ Private clinic   ☐ Other: ________'),
       ('State', ''), ('LGA', ''), ('Ward / town', ''), ('Street address or landmark', ''),
       ('Public phone number for patients (the facility line, NOT a personal number)', ''),
       ('Location (optional): Google Maps plus code, or latitude, longitude', '')], h=0.8)

para('', after=2)
para('2. Services (tick what the facility really offers today)', bold=True, size=10.5, color=NAVY, after=2)
table([('Antenatal care (ANC)', '☐ Yes   ☐ No'), ('Labour and delivery', '☐ Yes   ☐ No'),
       ('Emergency obstetric care (for example caesarean section, blood transfusion)', '☐ Yes, 24 hours   ☐ Yes, not 24 hours   ☐ No'),
       ('Postnatal care', '☐ Yes   ☐ No'), ('Newborn care', '☐ Yes   ☐ No'),
       ('Immunisation', '☐ Yes   ☐ No'), ('Family planning', '☐ Yes   ☐ No')], h=0.7)

para('', after=2)
para('3. Confirmation and permission', bold=True, size=10.5, color=NAVY, after=2)
table([('Details above were confirmed by (name, role)', ''),
       ('Date confirmed', ''),
       ('Permission: list this facility in the MamaCare AI directory (public)', '☐ Yes   ☐ No'),
       ('Permission: name this facility in the project record as a place where the app\'s danger-sign content was reviewed', '☐ Yes   ☐ No'),
       ('Officer-in-Charge or authorised person (name, role)', ''),
       ('Signature and date', '')], h=0.95)
para('', after=2)
para('We will correct or remove this facility at any time on request: najibadam@gmail.com. We will not share these details outside the app directory and the project record.', italic=True, size=9)
doc.save(out); print('saved', out)
